#!/usr/bin/env python3
"""Build the Multi-ERP Spend & Supplier Risk presales kit documents.

Mirrors the SAP 360 kit builders (Spend 360 in particular) so the kits read as a set:

    00_START_HERE.docx               what is in the kit and which file to open
    01_Management_Summary.docx       customer-safe summary
    02_Demo_Scripts_by_Persona.docx  one script per persona in the room
    03_SE_Quick_Start.docx           positioning, demo path, page tour, objections
    05_Architecture_and_Install.docx the layers, every object, and the build order
    06_Setup_and_Access.docx         public demo, local run, customer deploy, adapting

Every figure comes from /tmp/mfg_spend_facts.json, produced by tools/spend_risk_facts.py
against the running app server (which reads the live account). Nothing is transcribed by hand.

Three rules are enforced everywhere:
  * Say which time window a number belongs to: all PO history (dashboards), last four
    quarters (price variance) or trailing 12 months (Supplier Optimization).
  * All data is synthetic; third-party risk, ESG, sanctions, sub-tier and PPI are representative.
  * Issue dollars overlap, so they are never added up.

    python3 tools/spend_risk_facts.py --print && python3 tools/build_presales_kit.py
"""
from __future__ import annotations

import json
import pathlib
import sys
from datetime import date

from docx import Document
from docx.shared import Inches, Pt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

from docx_kit import (  # noqa: E402
    GREY,
    SAP_NAVY,
    SNOW_BLUE,
    body,
    bullet,
    callout,
    h1,
    h2,
    setup_page,
    table,
)

KIT = pathlib.Path.home() / "Documents" / "SAP" / "Multi_ERP_Spend_Presales_Kit"
SHOTS = KIT / "screenshots"
FACTS_FILE = pathlib.Path("/tmp/mfg_spend_facts.json")
DATE = date.today().strftime("%d %B %Y")
NAME = "Multi-ERP Spend & Supplier Risk"

PAGE_PURPOSE = {
    "ask-cortex": "Any question, answered by Cortex (AI_COMPLETE) grounded on the gold layer",
    "overview": "Global spend, golden suppliers, direct spend on contract, price-variance savings, high-risk spend",
    "suppliers": "Concentration, tail, and suppliers bought from in several ERPs",
    "price-variance": "Same material, same quarter: premium over the best plant, contract and PPI",
    "supplier-risk": "Composite risk tier, sanctions, ESG, sub-tier, region (representative third-party data)",
    "supplier-optimization": "Scenario levers, rule-based issues with Cortex remedy plans, Cortex recommendation",
    "contracts": "On- versus off-contract (maverick) spend by category, company and supplier",
    "match-payments": "PO vs goods receipt vs invoice exceptions, days to pay, lost early-payment discounts",
    "golden-record": "Cross-ERP supplier deduplication by tax ID, D-U-N-S, then name + country",
    "data-sources": "Choose the system behind each SAP instance, switch sources on and off, Apply runs dbt",
    "lineage": "The dbt DAG per canonical object, unselected adapters greyed out",
}

ISSUE_LABEL = {
    "MAVERICK_SPEND": "Maverick spend",
    "SUPPLIER_FRAGMENTATION": "Supplier fragmentation",
    "PAYMENT_TERMS_MISMATCH": "Payment-terms mismatch across ERPs",
    "SINGLE_SOURCE_SUBTIER": "Single-source sub-tier",
    "PLANT_PRICE_PREMIUM": "Plant price premium",
    "HIGH_RISK_CONCENTRATION": "High-risk concentration",
    "LOST_EARLY_PAY_DISCOUNT": "Lost early-pay discounts",
    "SANCTIONS_EXPOSURE": "Sanctions exposure",
}


def load_facts() -> dict:
    if not FACTS_FILE.exists():
        sys.exit(f"{FACTS_FILE} missing — run tools/spend_risk_facts.py --print first")
    return json.loads(FACTS_FILE.read_text())["facts"]


def usd(v) -> str:
    v = float(v)
    if abs(v) >= 1e9:
        return f"${v / 1e9:,.2f}B"
    if abs(v) >= 1e6:
        return f"${v / 1e6:,.1f}M"
    if abs(v) >= 1e3:
        return f"${v / 1e3:,.0f}K"
    return f"${v:,.0f}"


def pct(v, d=1) -> str:
    return f"{float(v):.{d}f}%"


def page(F, pid) -> str:
    return next(p["label"] for p in F["app_pages"] if p["id"] == pid)


def has_rec(F) -> bool:
    return bool(F.get("cortex_recommendation") and F["cortex_recommendation"].get("savings") is not None)


def shot(doc, name, caption, width=6.6):
    f = SHOTS / name
    if not f.exists():
        return
    doc.add_picture(str(f), width=Inches(width))
    body(doc, caption, size=8.5, italic=True, color=GREY)


def title_block(doc, title, subtitle, strap):
    for text_, size, bold, color, after in (
        (title, 20, True, SAP_NAVY, 2),
        (subtitle, 11.5, False, SNOW_BLUE, 2),
        (strap, 9, False, GREY, 14),
    ):
        p = doc.add_paragraph()
        r = p.add_run(text_)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(after)


def new_doc(subtitle, strap):
    doc = Document()
    setup_page(doc)
    title_block(doc, NAME, subtitle, strap)
    return doc


# ------------------------------------------------------------------ callouts

def synthetic_callout(doc, F):
    callout(
        doc,
        "All data is synthetic",
        "Meridian Industrial and every supplier are fictional. Purchase orders, receipts, invoices and payments "
        "are generated by data_gen/generate.py (seed 360) and rendered into each ERP's native tables. Supplier "
        "risk, sanctions, ESG, sub-tier and PPI data are representative, shaped like D&B, EcoVadis and Resilinc "
        "feeds but not real feeds. Say so when you open Supplier Risk.",
    )


def window_callout(doc, F):
    ob = F["opt_baseline"]
    callout(
        doc,
        "Three time windows — say which one",
        f"Dashboards use all PO history, Jan 2024 to Sep 2026 ({usd(F['total_spend_usd'])}). Price variance uses the "
        f"last four quarters, direct materials bought by two or more plants ({usd(F['price_comparable_spend_usd'])} "
        f"comparable). {page(F, 'supplier-optimization')} uses the trailing 12 months ({usd(ob['spend'])}). So "
        f"high-risk spend is {usd(F['high_risk_spend_usd'])} on the dashboard (all history) and "
        f"{usd(ob['high_risk_spend'])} in optimization (12 months). Both are right; quote the window.",
    )


def suppliers_callout(doc, F):
    g = F["golden"]
    callout(
        doc,
        "Suppliers: two counts",
        f"{g['golden']} golden suppliers come out of {g['source_records']} vendor-master records across the ERPs; "
        f"{F['suppliers_with_spend']} of them have PO spend. The other {g['golden'] - F['suppliers_with_spend']} "
        f"exist in a vendor master only. The dashboard shows {F['suppliers_with_spend']}.",
    )


def model_callout(doc, F):
    callout(
        doc,
        "Savings are a model, not a forecast",
        "The scenario model is deterministic: price harmonization is applied first, the other levers apply to the "
        "spend that remains, so nothing is double counted. Treat totals as directional. Cortex proposes lever "
        "values; the model, not Cortex, computes the savings. Issue dollars overlap — the same dollar can sit in "
        "several issues — so never add issue impacts up.",
    )


def sources_callout(doc, F):
    callout(
        doc,
        "Source switching reconciles — and shows its gaps",
        "Each SAP instance is generated once and rendered as ECC, S/4HANA, BW and BDC, so SAP spend is identical "
        f"whichever type is chosen ({usd(F['sap_spend_usd'])} of SAP spend; native PO value is identical across "
        f"{', '.join(F['sap_po_value_by_type'])}; enforced by the dbt test "
        "assert_sap_po_value_reconciles). The real gaps are shown, not hidden: BDC has no goods receipts or "
        "pricing conditions, BW has no pricing conditions, Coupa invoices are header-level (two-way match only), "
        "and the acquired ERPs carry no contracts.",
    )


def payments_callout(doc, F):
    callout(
        doc,
        "Late payments are generated",
        f"The synthetic payments run late {pct(F['payments']['late_pct'])} of the time. That is a property of the "
        f"generator, not a benchmark — do not quote it as one.",
    )


def agent_scope_callout(doc, F):
    sv = F["semantic_view"]
    callout(
        doc,
        "What the agent knows",
        f"{F['agent'].split('.')[-1]} (Snowflake Intelligence) answers from the semantic view {sv['name']} only: "
        f"{sv['tables']} tables, {sv['dimensions']} dimensions, {sv['facts']} facts, {sv['metrics']} metrics. It does "
        f"not see the app-side scenario simulation. In-app Ask Cortex is AI_COMPLETE grounded on the gold layer.",
    )


def public_callout(doc, F):
    callout(
        doc,
        "The public demo is a snapshot",
        f"{F['urls']['public_app']} is read-only: selection SAP A = ECC, SAP B = S/4HANA, all non-SAP sources on. "
        f"The scenario model runs in the browser. Cortex answers are pre-generated for the suggested questions, all "
        f"{F['issues_total']} remedies and one recommendation. Filters there combine an exact company selection, or "
        f"all/one company with one category. Data Sources cannot Apply.",
    )


def headline_rows(F):
    ob, od = F["opt_baseline"], F["opt_default"]
    return [
        ["Spend in scope", f"{usd(F['total_spend_usd'])} across {F['po_lines']:,} PO lines, all history Jan 2024–Sep 2026"],
        ["Source systems", f"{F['erps']} ERPs and P2P systems, {F['plants']} plants, one canonical model"],
        ["Supply base", f"{F['suppliers_with_spend']} golden suppliers with spend ({F['golden']['golden']} in total); top 10 hold {pct(F['top10_share_pct'])}"],
        ["Contract coverage", f"{pct(F['direct_on_contract_pct'])} of direct spend on contract; {usd(F['maverick_spend_usd'])} maverick (all history)"],
        ["Price variance", f"{usd(F['price_savings_usd'])} premium over best plant, {pct(F['price_savings_pct'])} of comparable spend (last 4 quarters)"],
        ["Risk (representative)", f"{F['high_risk_suppliers']} High/Critical suppliers, {usd(F['high_risk_spend_usd'])} spend (all history)"],
        ["Optimization (12 months)", f"Default levers save {usd(od['savings'])} ({pct(od['savings_pct'] * 100)}) of {usd(ob['spend'])}; {F['issues_total']} issues flagged"],
        ["Build", f"{F['bronze_tables']} bronze tables, {F['adapter_views']} dbt adapters, {F['dbt_counts']['models']} models, {len(F['app_pages'])}-page app"],
    ]


# ---------------------------------------------------------------- START HERE

def build_start_here(F):
    doc = new_doc("Presales kit — start here", f"Manufacturing use case 01  ·  {DATE}  ·  Owner: Dave Freriks")
    body(doc,
         "Multi-ERP Spend & Supplier Risk answers the manufacturing playbook question \"What is my global spend by "
         "supplier, and where do prices vary across plants?\" for a company running two SAP instances, three acquired "
         "non-SAP ERPs and Coupa. Each SAP instance can be read as ECC, S/4HANA, BW or SAP BDC. dbt adapters turn each "
         "system's native tables into one canonical model, gold marts and a semantic view feed a Cortex Agent and a "
         "React app. It is a portable, modular pattern — not tied to SAP BDC.")
    callout(doc, "Fastest path to a demo",
            f"Open the public demo at {F['urls']['public_app']} — no login, nothing to install. For live Cortex and "
            f"source switching, run the app locally ({F['urls']['local_app']}). See 06_Setup_and_Access.")
    synthetic_callout(doc, F)
    table(doc, ["The demo in eight numbers", "Value"], headline_rows(F), widths=[2.2, 4.5])
    body(doc,
         f"Every figure above was read from the running app (live account) when the kit was built on {DATE}, by "
         f"tools/spend_risk_facts.py, not copied from a previous deck.", italic=True)
    window_callout(doc, F)

    h1(doc, "Which file to open")
    table(doc, ["File", "Use it when", "Read time"], [
        ["03_SE_Quick_Start.docx", "You are demoing this week. Positioning, a ten-minute path, the agent "
         "questions, page tour, discovery questions, objections. Start here.", "10 min"],
        ["02_Demo_Scripts_by_Persona.docx", "You know who is in the room — CPO, category manager, supplier "
         "risk, CFO, architect. Each script opens on a different page.", "per script"],
        ["01_Management_Summary.docx", "Leaving something with the customer, or briefing an exec. "
         "Customer-safe and explicit about what is synthetic.", "15 min"],
        ["05_Architecture_and_Install.docx", "An architect is in the room, or you are standing it up yourself.", "reference"],
        ["06_Setup_and_Access.docx", "You need access: public demo, local run, deploying into a customer account.", "5 min"],
        ["00_Presales_Overview.pptx", "Introducing the pattern to a customer or partner: the problem, the portable architecture, "
         "the demo in numbers.", "10 slides"],
        ["Multi_ERP_Spend_Demo.pptx", "Demoing without a live app, or leaving a screen-by-screen walkthrough behind.", "1 per page"],
        ["Multi_ERP_Spend_Walkthrough.mp4", "Watching the narrated demo end to end before you give it.", "video"],
        ["screenshots/", "Every page, the Supplier Optimization tabs and a Cortex remedy, at 2x resolution.", "images"],
    ], widths=[2.1, 3.8, 0.8])
    model_callout(doc, F)
    suppliers_callout(doc, F)

    h1(doc, "The companion kits")
    body(doc,
         "The SAP BDC 360 kits (Finance, Spend, Supply Chain, People, Sales, Cloud ERP) follow the same document "
         "pattern. Spend 360 shows one SAP BDC landscape; this use case shows the mixed landscape most manufacturers "
         "actually run: ECC and S/4 side by side, acquisitions on other ERPs, and an indirect P2P tool.")
    h1(doc, "Support")
    body(doc, f"Source and documentation: {F['urls']['repo']} (use_cases/01_multi_erp_spend). Questions: Dave Freriks.")
    out = KIT / "00_START_HERE.docx"
    doc.save(out)
    return out


# ---------------------------------------------------------- MANAGEMENT SUMMARY

def build_management_summary(F):
    doc = new_doc("Management summary", f"Spend and supplier risk across every ERP  ·  {DATE}")
    h1(doc, "The problem")
    body(doc,
         "Manufacturers rarely run one ERP. Two SAP instances on different releases, acquired companies on Oracle, "
         "Dynamics or NetSuite, and indirect spend in Coupa: the same supplier appears under several names, the same "
         "material is bought at different prices in different plants, and supplier risk is assessed apart from spend. "
         "Answering \"what do we spend with whom, and where do prices vary\" means a consolidation project every time.")
    h1(doc, "What it does")
    for t in [
        "Lands each system's tables unchanged and maps them with dbt adapters into one canonical model — SAP ECC, S/4HANA, BW or BDC per instance, plus Oracle EBS, Dynamics 365, NetSuite and Coupa.",
        "Builds a golden supplier record across ERPs (tax ID, D-U-N-S, then name and country).",
        "Compares prices for the same material across plants, and joins spend to supplier risk, sanctions, ESG and sub-tier exposure.",
        "Models savings scenarios, flags rule-based issues, and uses Cortex to draft remedy plans and recommend a scenario.",
        "Lets anyone ask questions in plain English through a Cortex Agent over a governed semantic view.",
    ]:
        bullet(doc, t)

    h1(doc, "What the demo dataset shows")
    table(doc, ["Measure", "Value"], headline_rows(F)[:7], widths=[2.2, 4.5])
    synthetic_callout(doc, F)

    h2(doc, "Supplier risk tiers (all history)")
    table(doc, ["Tier", "Suppliers", "Spend"],
          [[t["tier"], str(t["suppliers"]), usd(t["spend"])] for t in F["risk_tiers"]], widths=[2.8, 1.5, 2.4])

    h2(doc, "Savings scenario (trailing 12 months)")
    lv = F["opt_default_by_lever"]
    rows = [["Price harmonization", usd(lv["price"]), usd(F["opt_full_potential"]["price"])],
            ["Maverick to contract", usd(lv["maverick"]), usd(F["opt_full_potential"]["maverick"])],
            ["Tail consolidation", usd(lv["tail"]), usd(F["opt_full_potential"]["tail"])],
            ["Early-pay discounts", usd(lv["discounts"]), usd(F["opt_full_potential"]["discounts"])],
            ["Risk shift (switching premium)", usd(lv["switching"]), "—"],
            ["Total", usd(F["opt_default"]["savings"]), "—"]]
    table(doc, ["Lever", "Default levers", "Full potential"], rows, widths=[2.8, 1.9, 2.0])
    txt = (f"Baseline {usd(F['opt_baseline']['spend'])} over 12 months. ")
    if has_rec(F):
        r = F["cortex_recommendation"]
        txt += (f"Cortex's recommended lever set models {usd(r['savings'])} ({pct(r['savings_pct'] * 100)}). ")
    txt += "Directional: a deterministic model applied in sequence, not a forecast."
    body(doc, txt, italic=True)

    h1(doc, "What is real, and what is representative")
    window_callout(doc, F)
    model_callout(doc, F)

    h1(doc, "Why Snowflake")
    for t in [
        "One platform for every source: SAP and non-SAP land side by side, and adapters — not a migration — make them comparable.",
        "dbt runs inside Snowflake (EXECUTE DBT PROJECT), so the pipeline is versioned in Git and governed in the account.",
        "The source behind each SAP instance can change (ECC today, S/4HANA or BDC tomorrow) without rebuilding the analytics.",
        "One semantic view serves the dashboards, the agent and any BI tool — one definition of spend everywhere.",
        "Cortex works on the governed data: remedy plans and scenario recommendations without data leaving the account.",
    ]:
        bullet(doc, t)
    h1(doc, "Suggested next step")
    body(doc,
         "A scoped proof of value on the customer's own systems: land two sources unchanged (for example one SAP "
         "instance and one acquired ERP), adapt the staging adapters to their Z-fields and company codes, and run the "
         "same canonical model, reconciliation tests, semantic view and agent over them.")
    out = KIT / "01_Management_Summary.docx"
    doc.save(out)
    return out


# -------------------------------------------------------------- DEMO SCRIPTS

def personas(F):
    P = lambda pid: page(F, pid)  # noqa: E731
    ob, od = F["opt_baseline"], F["opt_default"]
    aps = F["ask_page_suggestions"]
    ms = {m["status"]: m for m in F["match_status"]}
    top = F["top_issues"]
    rec = F["cortex_recommendation"] if has_rec(F) else None
    return [
        ("CPO / VP Procurement", "What do we spend with whom across every ERP, and how much of it do we control?",
         P("overview"), [
             (P("overview"), f"{usd(F['total_spend_usd'])} across {F['erps']} systems and {F['plants']} plants, all history. "
                             f"{pct(F['direct_on_contract_pct'])} of direct spend on contract."),
             (P("suppliers"), f"Top 10 hold {pct(F['top10_share_pct'])}; {F['suppliers_for_80pct']} suppliers make 80% of spend. Show suppliers bought in several ERPs."),
             (P("contracts"), f"{usd(F['maverick_spend_usd'])} maverick spend (all history). That is the controllable number."),
             (P("price-variance"), f"{usd(F['price_savings_usd'])} paid above the best plant for the same material, last four quarters."),
             (P("supplier-optimization"), f"Default levers save {usd(od['savings'])} of {usd(ob['spend'])} (trailing 12 months). Move one slider."),
             (P("supplier-optimization") + " — Cortex Recommendation",
              (f"Cortex proposes levers worth {usd(rec['savings'])}; the model computes the total. Apply to builder." if rec
               else "Generate a Cortex recommendation; the model, not Cortex, computes the total. Apply to builder.")),
             (P("ask-cortex"), f"Ask: '{F['ask_suggestions'][0]}'"),
         ], "One view of spend across every ERP, without waiting for the S/4 migration or the acquisitions to be integrated."),
        ("Category Manager (direct materials)", "Where do plants pay different prices, and where can I consolidate?",
         P("price-variance"), [
             (P("price-variance"), f"{F['price_materials']} materials bought by two or more plants; {usd(F['price_savings_usd'])} premium "
                                   f"({pct(F['price_savings_pct'])} of {usd(F['price_comparable_spend_usd'])}), last four quarters. Show the PPI line."),
             (P("golden-record"), f"{F['golden']['source_records']} vendor records collapse to {F['golden']['golden']} golden suppliers; "
                                  f"{F['golden']['multi_source']} appear in more than one system. That is what makes the comparison fair."),
             (P("suppliers"), f"{F['suppliers_for_80pct']} suppliers make 80% of spend. The tail is the consolidation target."),
             (P("contracts"), "Filter to your category: off-contract spend by supplier and company."),
             (P("supplier-optimization") + " — Issues & Remedies",
              f"Open a {ISSUE_LABEL['SUPPLIER_FRAGMENTATION'].lower()} or {ISSUE_LABEL['PLANT_PRICE_PREMIUM'].lower()} issue; generate the Cortex remedy plan."),
             (P("ask-cortex"), f"Ask: '{aps['price-variance'][0]}'"),
         ], "Price harmonization is the first lever because the evidence is the company's own purchase history."),
        ("Supplier Risk & Compliance lead", "Which suppliers could hurt us, and how much do we spend with them?",
         P("supplier-risk"), [
             (P("supplier-risk"), f"{F['high_risk_suppliers']} High/Critical suppliers carry {usd(F['high_risk_spend_usd'])} (all history). "
                                  "Say the risk data is representative."),
             (P("supplier-risk"), f"{F['risk_coverage']['dnb_matched']} of {F['risk_coverage']['suppliers']} matched to the risk feed; "
                                  f"{F['risk_coverage']['single_source']} single-source; {F['risk_coverage']['sanctions_hits']} sanctions hit."),
             (P("supplier-optimization") + " — Issues & Remedies",
              f"{F['issues_total']} issues ({F['issues_by_severity'].get('Critical', 0)} Critical). Open the top one: {top[0]['entity']} — {top[0]['summary']}."),
             (P("supplier-optimization") + " — Cortex remedy", "Generate the remedy: diagnosis, ordered steps with owners, ERP actions, KPIs."),
             (P("supplier-optimization") + " — Scenario Builder",
              f"Raise the risk-shift lever: High/Critical spend {usd(ob['high_risk_spend'])} → {usd(od['high_risk_spend'])} at defaults (12 months)."),
             (P("golden-record"), "One supplier, one risk score, across every ERP that buys from it."),
             (P("ask-cortex"), f"Ask: '{aps['supplier-risk'][1]}'"),
         ], "Risk ranked by exposure — spend across every ERP — not by an alphabetical vendor list."),
        ("CFO / Finance", "Where is cash leaking: discounts, payment terms, maverick spend?",
         P("match-payments"), [
             (P("match-payments"), f"{usd(F['payments']['paid'])} paid, average {F['payments']['avg_days']:.0f} days to pay. "
                                   f"{usd(F['payments']['discount_lost'])} early-pay discount lost vs {usd(F['payments']['discount_taken'])} taken."),
             (P("match-payments"), f"Exceptions: {usd(ms['Price variance']['spend'])} price variance, {usd(ms['Invoice without receipt']['spend'])} "
                                   f"invoiced without receipt. Coupa is two-way only ({usd(ms['Two-way only (header invoice)']['spend'])})."),
             (P("contracts"), f"{usd(F['maverick_spend_usd'])} maverick spend (all history) — negotiated savings leak here."),
             (P("supplier-optimization") + " — Issues & Remedies",
              f"{ISSUE_LABEL['PAYMENT_TERMS_MISMATCH']}: the same supplier on different terms in different ERPs."),
             (P("supplier-optimization") + " — Scenario Builder",
              f"Default levers: {usd(od['savings'])} over 12 months; on-contract {pct(ob['on_contract_pct'] * 100)} → {pct(od['on_contract_pct'] * 100)}."),
             (P("ask-cortex"), f"Ask: '{F['ask_suggestions'][4]}'"),
         ], "Working capital and savings from data finance already owns — no new system, no new extract."),
        ("Enterprise / Data Architect", "How is this built, and how does it cope with ECC → S/4, BDC and acquisitions?",
         P("data-sources"), [
             (P("data-sources"), f"{len(F['source_slots'])} source slots. SAP A and B can each be ECC, S/4HANA, BW or BDC. Preview the native tables."),
             (P("data-sources"), "Switch SAP B to BDC and Apply: dbt build runs (locally, or EXECUTE DBT PROJECT in Snowflake). Coverage shows BDC has no goods receipts or pricing conditions."),
             (P("lineage"), f"The dbt DAG: {F['adapter_views']} adapters → canonical silver → gold; unselected adapters greyed out."),
             (P("overview"), "SAP spend is unchanged after the switch — the reconciliation test enforces it."),
             (P("golden-record"), f"Match methods: " + ", ".join(f"{m['method']} {m['records']}" for m in F["match_methods"]) + "."),
             (P("ask-cortex"), f"Ask: '{F['ask_suggestions'][5]}'; for generated SQL, ask {F['agent'].split('.')[-1]} in Snowflake Intelligence."),
         ], "The S/4 migration, BDC and the next acquisition become a change of adapter, not a rebuild of the analytics."),
    ]


def build_demo_scripts(F):
    doc = new_doc("Demo scripts by persona", f"Five scripts, each opening on a different page  ·  {DATE}")
    body(doc, "Pick the script for the most senior person in the room. Each runs five to eight minutes.")
    window_callout(doc, F)
    for name, question, opener, steps, close in personas(F):
        h1(doc, name)
        body(doc, f"Their question: {question}", italic=True)
        body(doc, f"Open on: {opener}")
        table(doc, ["#", "Page", "Do and say"],
              [[str(i), p, s] for i, (p, s) in enumerate(steps, 1)], widths=[0.3, 1.9, 4.5])
        callout(doc, "Close with", close)
    out = KIT / "02_Demo_Scripts_by_Persona.docx"
    doc.save(out)
    return out


# --------------------------------------------------------------- QUICK START

SHOT_FOR = {p: f"{i:02d}_{p}.png" for i, p in enumerate(PAGE_PURPOSE)}


def build_quick_start(F):
    P = lambda pid: page(F, pid)  # noqa: E731
    doc = new_doc("SE quick start", f"Ten-minute demo path, the agent's scope, and objection handling  ·  {DATE}")
    body(doc, "Read once, the day before you demo. It assumes you have not opened the app.")
    h1(doc, "Positioning, in three sentences")
    bullet(doc, "Manufacturers run several ERPs — SAP on two releases, acquisitions on others — and cannot see spend, "
                "prices and supplier risk across them without a consolidation project.")
    bullet(doc, "Landing each system unchanged in Snowflake and mapping it with dbt adapters gives one canonical model, "
                "whatever mix of ECC, S/4HANA, BW, SAP BDC and non-SAP ERPs the customer runs.")
    bullet(doc, "Snowflake adds the governed semantic view, Cortex for remedies and recommendations, and the app.")
    h1(doc, "Before you start")
    bullet(doc, f"Rehearse on the public demo ({F['urls']['public_app']}). For live Cortex or source switching, run locally.")
    bullet(doc, "Know the three windows (below) and quote the window with every number.")
    bullet(doc, "Decide how you will introduce the synthetic data. Say it in the first minute, and again on Supplier Risk.")
    synthetic_callout(doc, F)
    window_callout(doc, F)

    h1(doc, "The ten-minute path")
    od, ob = F["opt_default"], F["opt_baseline"]
    table(doc, ["#", "Page", "Do and say", "Time"], [
        ["1", P("overview"), f"{usd(F['total_spend_usd'])} across {F['erps']} systems, all history. {pct(F['direct_on_contract_pct'])} of direct spend on contract.", "1 min"],
        ["2", P("data-sources"), "Six slots; SAP A and B each ECC, S/4HANA, BW or BDC. Show the coverage grid.", "1 min"],
        ["3", P("golden-record"), f"{F['golden']['source_records']} vendor records → {F['golden']['golden']} golden suppliers.", "1 min"],
        ["4", P("price-variance"), f"{usd(F['price_savings_usd'])} above best plant, last four quarters.", "1 min"],
        ["5", P("supplier-risk"), f"{F['high_risk_suppliers']} High/Critical suppliers, {usd(F['high_risk_spend_usd'])} (all history). Representative.", "1 min"],
        ["6", P("supplier-optimization"), f"Scenario: {usd(od['savings'])} of {usd(ob['spend'])} at defaults (12 months). Open an issue, generate a remedy.", "2 min"],
        ["7", P("ask-cortex"), "Ask a suggested question live; show governed SQL via the agent in Snowflake Intelligence.", "2 min"],
        ["8", P("lineage"), "Close here for a technical room: adapters → silver → gold.", "1 min"],
    ], widths=[0.3, 1.6, 4.1, 0.7])

    h2(doc, "The questions the app ships with")
    body(doc, "Read from the app's source, so these are exactly the chips on screen (Ask Cortex).")
    for q in F["ask_suggestions"]:
        bullet(doc, q)
    agent_scope_callout(doc, F)

    h1(doc, "What is on each page")
    table(doc, ["Page", "What it shows"], [[p["label"], PAGE_PURPOSE.get(p["id"], "—")] for p in F["app_pages"]],
          widths=[2.0, 4.7])

    h1(doc, "Page tour")
    body(doc, "Screenshots from the running app, in navigation order.")
    for p in F["app_pages"]:
        h2(doc, p["label"])
        shot(doc, SHOT_FOR.get(p["id"], ""), PAGE_PURPOSE.get(p["id"], ""))
        if p["id"] == "supplier-optimization":
            shot(doc, "x1_issues.png", f"Issues & Remedies: {F['issues_total']} rule-based issues. Do not add the dollars up — they overlap.")
            shot(doc, "x2_remedy.png", "A Cortex remedy plan: diagnosis, ordered steps, ERP actions, KPIs.")
            shot(doc, "x3_cortex_rec.png", "Cortex Recommendation: proposed levers; the scenario model computes the total.")
        if p["id"] == "data-sources":
            shot(doc, "x4_sources_bdc.png", "SAP BDC chosen for an instance: coverage shows no goods receipts or pricing conditions.")
        if p["id"] == "ask-cortex":
            shot(doc, "x5_ask_cortex.png", "An Ask Cortex answer to a suggested question.")

    h1(doc, "Issues by type (trailing 12 months)")
    table(doc, ["Issue type", "Issues", "Dollars at stake"],
          [[ISSUE_LABEL.get(x["issue_type"], x["issue_type"]), str(x["issues"]), usd(x["impact_usd"])] for x in F["issues_by_type"]],
          widths=[3.2, 1.2, 2.3])
    body(doc, "Dollars overlap across issues; show them per issue, never as a total.", italic=True)

    h1(doc, "Discovery questions")
    for q in [
        "How many ERPs and P2P tools hold purchasing data today, and which came with acquisitions?",
        "Where are you on ECC to S/4HANA — and will both run side by side for a while?",
        "Is SAP Business Data Cloud in your plans, or do you replicate SAP tables (SLT, Openflow, a partner)?",
        "How long does it take to answer 'what do we spend with this supplier across all companies'?",
        "Do you compare prices for the same material across plants today? Who acts on it?",
        "Where does supplier risk data come from (D&B, EcoVadis, Resilinc), and does it connect to spend?",
        "Do you run dbt today, and where?",
    ]:
        bullet(doc, q)

    h1(doc, "Objections, and what to say")
    table(doc, ["They say", "You say"], [
        ["Is this real data?", "No — all synthetic, generated with a fixed seed, and the company and suppliers are fictional. "
         "Risk, ESG, sanctions, sub-tier and PPI data are representative of third-party feeds."],
        ["Why does high-risk spend differ between pages?", f"Different windows: {usd(F['high_risk_spend_usd'])} is all history on the "
         f"dashboard; {usd(ob['high_risk_spend'])} is the trailing 12 months in Supplier Optimization."],
        ["Is the savings number a forecast?", "No. It is a deterministic model with levers applied in sequence. Cortex proposes "
         "levers; the model computes totals. Directional, for prioritising."],
        ["We are moving to S/4 / BDC — does this break?", "No. Switch the instance on Data Sources: SAP spend reconciles exactly, "
         "and the coverage view shows what BDC does not carry (goods receipts, pricing conditions)."],
        ["Our ERPs have custom Z-fields.", "Adapters are generated from tools/build_silver_sql.py; add the fields there and rerun."],
        ["Can the agent see my scenario?", "No. It answers from the semantic view; the scenario model lives in the app."],
        ["Why is late payment so high?", f"{pct(F['payments']['late_pct'])} is a property of the synthetic generator, not a benchmark."],
    ], widths=[1.9, 4.8])

    h1(doc, "If it goes wrong")
    bullet(doc, "A page is empty — check the company and category filters; they persist across pages.")
    bullet(doc, "Ask Cortex is slow — AI_COMPLETE takes several seconds; ask a suggested question on the public demo for an instant answer.")
    bullet(doc, "A number disagrees with this document — check the source selection on Data Sources (the kit used SAP A = ECC, "
                "SAP B = S/4HANA, all non-SAP on); regenerate the kit from the facts script.")
    out = KIT / "03_SE_Quick_Start.docx"
    doc.save(out)
    return out


# ------------------------------------------------- ARCHITECTURE AND INSTALL

def build_architecture(F):
    doc = new_doc("Architecture and install", f"Bronze to agent, every object, and the build order  ·  {DATE}")
    sv = F["semantic_view"]
    body(doc, "Everything sits in the MFG_SPEND_RISK database. Source tables land unchanged; dbt adapters map them into a "
              f"canonical model. Object counts below were read from the account on {DATE}.")
    h1(doc, "The stack")
    table(doc, ["Layer", "Objects", "What it does"], [
        ["Bronze", f"{F['bronze_tables']} tables in {F['bronze_schemas']} SRC_* schemas, {F['bronze_rows']:,} rows",
         "Native tables per system and instance: ECC, S/4HANA, BW, BDC, Oracle EBS, Dynamics 365, NetSuite, Coupa, third-party."],
        ["dbt staging", f"{F['adapter_views']} adapter views in ADAPTERS", "One per source system × instance × canonical object; generated by tools/build_silver_sql.py."],
        ["Silver", "Canonical objects in SILVER", "union_selected macro unions only the selected adapters; golden-supplier cross-reference; AI_CLASSIFY category map."],
        ["Gold", f"{len(F['gold_dynamic_tables'])} dynamic tables, {len(F['gold_tables'])} tables in GOLD",
         ", ".join(F["gold_dynamic_tables"]) + "; " + ", ".join(F["gold_tables"]) + "."],
        ["Semantic", sv["name"], f"{sv['tables']} tables, {sv['dimensions']} dimensions, {sv['facts']} facts, {sv['metrics']} metrics, {sv['relationships']} relationships."],
        ["Agent", F["agent"], "Cortex Agent with Cortex Analyst over the semantic view, in Snowflake Intelligence."],
        ["App", f"{len(F['app_pages'])} pages", "React client (5188) and Express server (3013); Cortex via AI_COMPLETE."],
    ], widths=[0.9, 2.2, 3.6])

    h1(doc, "Source slots")
    table(doc, ["Slot", "Scope", "Allowed systems", "Kit selection"],
          [[s["label"], s["scope"], ", ".join(s["allowed_systems"]), s["system_code"] if s["enabled"] else "off"] for s in F["source_slots"]],
          widths=[1.8, 2.4, 1.5, 1.0])
    sources_callout(doc, F)

    h1(doc, "How source selection works")
    body(doc, "Data Sources writes CONFIG.SOURCE_SELECTION and runs dbt build --full-refresh with the selection as a var. "
              "Silver is rebuilt from the selected adapters, gold dynamic tables are replaced, and the tests run on the result. "
              "Runs are logged in CONFIG.REBUILD_LOG.")
    lr = F["last_rebuild"]
    table(doc, ["dbt", "Detail"], [
        ["dbt project object", f"{F['dbt_project']['name']} (dbt {F['dbt_project']['dbt_version']})"],
        ["Git repository object", F["git_repository"]],
        ["Project size", f"{F['dbt_counts']['models']} models, {F['dbt_counts']['tests']} tests, {F['dbt_counts']['sources']} sources"],
        ["Last source-switch rebuild (Apply)", f"{lr['finished_at'][:19]} — {lr['status']}, {lr['message']}"],
        ["Execution", "DBT_EXECUTION=local runs the dbt CLI; DBT_EXECUTION=snowflake runs EXECUTE DBT PROJECT"],
    ], widths=[1.8, 4.9])

    h1(doc, "Supplier Optimization")
    lv = F["levers"]
    table(doc, ["Lever", "Range", "Default"],
          [[v["label"], f"{v['min']}–{v['max']}", str(v["default"])] for v in lv.values()], widths=[3.4, 1.6, 1.7])
    body(doc, "app/server/src/routes/optimize.ts applies the levers in sequence over GOLD.MART_SUPPLIER_OPPORTUNITY — price "
              "first, the others on the remaining spend. Issues come from GOLD.MART_SUPPLIER_ISSUES; Cortex remedies and the "
              "recommendation use AI_COMPLETE. Scenarios are saved in CONFIG.SUPPLIER_SCENARIOS.")
    model_callout(doc, F)

    h1(doc, "Build order")
    table(doc, ["Step", "What to run", "Result"], [
        ["0", "Repo root: python3.13 -m venv .venv; pip install dbt-core~=1.12 dbt-snowflake~=1.12", "Shared dbt environment"],
        ["1", "data_gen/generate.py, data_gen/load_bronze.py", "MFG_SPEND_RISK and the SRC_* bronze schemas"],
        ["2", "sql/03_config.sql", "CONFIG: source slots, selection, logs"],
        ["3", "tools/build_silver_sql.py", "dbt staging adapters, sources.yml, docs/source_mapping.md"],
        ["4", "dbt build (local) or snowflake/setup_git_dbt.sql then EXECUTE DBT PROJECT", "Silver and gold"],
        ["5", "sql/04_semantic_view.sql", sv["name"]],
        ["6", "sql/05_cortex_agent.sql", F["agent"]],
        ["7", "app: cp .env.example .env; npm install; npm run dev", F["urls"]["local_app"]],
    ], widths=[0.5, 3.4, 2.8])
    body(doc, "After step 6 you can demo through Snowflake Intelligence without the app.")
    agent_scope_callout(doc, F)
    out = KIT / "05_Architecture_and_Install.docx"
    doc.save(out)
    return out


# ------------------------------------------------------------- SETUP, ACCESS

def build_setup(F):
    doc = new_doc("Setup and access", f"Public demo, local run, customer deploy  ·  {DATE}")
    h1(doc, "How to get to it")
    table(doc, ["Route", "What you get", "Where"], [
        ["Public static demo", "Credential-free browser build: all pages, scenario model in the browser, pre-generated Cortex answers.",
         F["urls"]["public_app"]],
        ["Local app", "The live app on the demo account: free-text Ask Cortex, live remedies, source switching with dbt.", F["urls"]["local_app"]],
        ["Snowflake Intelligence", "The agent over the semantic view, with generated SQL.", F["agent"]],
        ["Source repository", "Data generator, dbt project, SQL, app and these kit builders. Public.", F["urls"]["repo"]],
    ], widths=[1.4, 3.0, 2.3])
    callout(doc, "No Native App", "Unlike the SAP BDC 360 apps there is no Native App or Marketplace listing for this use case. "
                                  "Use the public demo, run it locally, or deploy it into an account.")
    public_callout(doc, F)

    h1(doc, "Running locally")
    for t in [
        "Node.js (npm) for the app; Python 3.13 for the generator and dbt.",
        "Repo-root virtualenv .venv with dbt-core and dbt-snowflake 1.12 (dbt does not run on Python 3.14 yet).",
        "A key-pair connection in ~/.snowflake/connections.toml; SNOWFLAKE_CONNECTION selects it (default dfreriksdemo).",
        "cp app/.env.example app/.env, fill in account, user and key path; cd app && npm install && npm run dev.",
        "Server on :3013, client on :5188. Set DBT_EXECUTION=snowflake to make Apply run the dbt project object.",
    ]:
        bullet(doc, t)

    h1(doc, "Deploying into a customer account")
    table(doc, ["Step", "What to run"], [
        ["1", "Generate and load bronze: data_gen/generate.py, data_gen/load_bronze.py"],
        ["2", "sql/03_config.sql (via framework/tools/sf.py)"],
        ["3", "snowflake/setup_git_dbt.sql once: API integration, Git repository object, dbt project object"],
        ["4", "EXECUTE DBT PROJECT ... ARGS = 'build --full-refresh --exclude staging int_category_map'"],
        ["5", "sql/04_semantic_view.sql, sql/05_cortex_agent.sql"],
        ["6", "Run the app against the account, or demo through Snowflake Intelligence"],
    ], widths=[0.5, 6.2])

    h1(doc, "Adapting to a customer's systems")
    for t in [
        "Land the customer's tables unchanged in a new SRC_<SYSTEM>_<SLOT> schema (Openflow, SLT or a partner for ECC/S4; an SAP BDC share; BW open hub or ODP; Fivetran or Openflow for Oracle, Dynamics, NetSuite; the Coupa API).",
        "Add or adjust the adapter in tools/build_silver_sql.py (custom Z-fields, company codes) and rerun it.",
        "Register the slot in CONFIG.SOURCE_SLOT and var('selection').",
        "dbt build: the reconciliation test, canonical column contract and coverage mart show what the new source supplies.",
    ]:
        bullet(doc, t)

    h1(doc, "Snowflake objects")
    table(doc, ["Object", "Detail"], [
        ["Database", "MFG_SPEND_RISK"],
        ["Bronze", f"{F['bronze_schemas']} SRC_* schemas, {F['bronze_tables']} tables"],
        ["Gold", ", ".join(F["gold_dynamic_tables"] + F["gold_tables"])],
        ["Semantic view", F["semantic_view"]["name"]],
        ["Agent", F["agent"]],
        ["dbt project", F["dbt_project"]["name"]],
        ["Git repository", F["git_repository"]],
        ["Demo connection", "dfreriksdemo"],
    ], widths=[1.6, 5.1])

    h1(doc, "Related assets")
    table(doc, ["Asset", "Where"], [
        ["SAP Spend 360 kit", "Single SAP BDC landscape, same document pattern"],
        ["Framework", "framework/docs/FRAMEWORK.md in the repo — the pattern for further mfg-* use cases"],
        ["Source mapping", "use_cases/01_multi_erp_spend/docs/source_mapping.md (generated)"],
    ], widths=[2.2, 4.5])
    h1(doc, "Contact")
    body(doc, "Dave Freriks — for access help, or to scope pointing this at a customer's own ERPs.")
    out = KIT / "06_Setup_and_Access.docx"
    doc.save(out)
    return out


def main() -> int:
    F = load_facts()
    KIT.mkdir(parents=True, exist_ok=True)
    for fn in (build_start_here, build_management_summary, build_demo_scripts, build_quick_start,
               build_architecture, build_setup):
        print(f"wrote {fn(F)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
